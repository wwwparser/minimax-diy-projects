"""Шаг 4. Сметы по каталогу Петровича: data/projects.json -> out/projects.xlsx, out/estimates.json, out/PROJECTS.md.

Позиция сметы: {"code": артикул Петровича, "qty": кол-во в единицах каталога, "kind": материал|крепёж|отделка|оснастка|расходник, "for": "зачем"}.
Цена берётся из локальной копии каталога (розница, Москва, дата сбора — 16.09.2026).
"""
import json, sqlite3, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill

ROOT = Path(__file__).resolve().parent.parent
BUDGET = 3000
KIND_ORDER = ["материал", "крепёж", "фурнитура", "отделка", "оснастка", "расходник"]


def engineering_override(p):
    """Если проект прошёл diy-project-engineer: текст из engineering/corrections.yaml, смета — закупка движка."""
    import yaml
    corr = (yaml.safe_load((ROOT / "engineering" / "corrections.yaml").read_text(encoding="utf-8")) or {}).get(p["id"], {})
    src = next(iter(sorted((ROOT / "engineering" / "build").glob(f"{p['id']:02d}_*/html_data.json"))), None)
    if not src:
        return None
    d = json.loads(src.read_text(encoding="utf-8"))
    lines = []
    for l in d["bom_purchase"]:
        if l.get("sum") is None:
            continue
        lines.append({"code": int(l["sku"]) if str(l.get("sku") or "").isdigit() else l.get("sku"), "qty": l["qty"],
                      "kind": l["kind"], "for": l.get("note") or "", "title": l["name"], "price": l["price"],
                      "unit": l.get("unit") or "шт", "sum": l["sum"], "url": l.get("url") or ""})
    lines.sort(key=lambda l: KIND_ORDER.index(l["kind"]) if l["kind"] in KIND_ORDER else 9)
    q = {k: v for k, v in p.items() if k != "bom"}
    for k in ("size", "summary", "steps"):
        if corr.get(k):
            q[k] = corr[k]
    q.update(bom=lines, total=round(d["total"]), tools_total=round(d["tools_total"]), engineered=True,
             fixes=corr.get("fixes", []))
    return q

def main():
    con = sqlite3.connect(ROOT / "data" / "catalog.sqlite3")
    projects = json.loads((ROOT / "data" / "projects.json").read_text(encoding="utf-8"))
    out, problems = [], []
    for p in projects:
        eng = engineering_override(p)
        if eng:
            out.append(eng)
            continue
        lines, total = [], 0
        for it in p["bom"]:
            row = con.execute("select title,price,unit,url,image from products where code=?", (it["code"],)).fetchone()
            if not row or not row[1]:
                problems.append(f'{p["id"]}: нет артикула {it["code"]}'); continue
            title, price, unit, url, image = row
            s = round(price * it["qty"], 2); total += s
            lines.append({**it, "title": title, "price": price, "unit": unit, "sum": s, "url": url})
        lines.sort(key=lambda l: KIND_ORDER.index(l["kind"]))
        p2 = {**{k: v for k, v in p.items() if k != "bom"}, "bom": lines, "total": round(total),
              "tools_total": round(sum(l["sum"] for l in lines if l["kind"] == "оснастка"))}
        if total > BUDGET:
            problems.append(f'{p["id"]} {p["name"]}: {total:.0f} ₽ > {BUDGET}')
        out.append(p2)
    (ROOT / "out").mkdir(exist_ok=True)
    (ROOT / "out" / "estimates.json").write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    write_xlsx(out); write_md(out)
    print(f"проектов: {len(out)}; дороже {BUDGET}: {sum(p['total'] > BUDGET for p in out)}")
    print("\n".join(problems) or "проблем нет")
    for p in sorted(out, key=lambda p: p["total"]):
        print(f'{p["id"]:>3} {p["total"]:>5} ₽  {p["material"]:<14} {p["name"]}')

def write_xlsx(projects):
    wb = Workbook(); ws = wb.active; ws.title = "Проекты"
    head = ["№", "Проект", "Материал", "Сложность", "Время, ч", "Итого, ₽", "в т.ч. оснастка, ₽", "Топ-10", "Что получится", "Источник"]
    ws.append(head)
    for p in projects:
        ws.append([p["id"], p["name"], p["material"], p["difficulty"], p["hours"], p["total"], p["tools_total"],
                   "да" if p.get("top") else "", p["summary"], p["source"]])
    b = wb.create_sheet("Сметы")
    b.append(["№", "Проект", "Тип", "Позиция Петровича", "Артикул", "Кол-во", "Ед.", "Цена, ₽", "Сумма, ₽", "Зачем", "Ссылка"])
    for p in projects:
        for l in p["bom"]:
            b.append([p["id"], p["name"], l["kind"], l["title"], l["code"], l["qty"], l["unit"], l["price"], l["sum"], l.get("for", ""), l["url"]])
        b.append([p["id"], p["name"], "ИТОГО", "", "", "", "", "", p["total"]]); b.append([])
    for sh in (ws, b):
        for c in sh[1]:
            c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5D50")
        sh.freeze_panes = "A2"
    for col, w in zip("ABCDEFGHIJ", [5, 40, 14, 10, 9, 10, 12, 8, 70, 50]): ws.column_dimensions[col].width = w
    for col, w in zip("ABCDEFGHIJK", [5, 34, 11, 70, 10, 7, 6, 9, 10, 40, 45]): b.column_dimensions[col].width = w
    for row in b.iter_rows(min_row=2):
        if row[2].value == "ИТОГО":
            for c in row: c.font = Font(bold=True)
    wb.save(ROOT / "out" / "projects.xlsx")

def write_md(projects):
    L = ["# Проекты под шуруповёрт Интерскол ДА-10/12В (36 Н·м)", "",
         "Цены — розница Петровича (Москва), снимок каталога 16.09.2026. В смету включено всё, что надо купить, если дома есть только шуруповёрт.", ""]
    for p in sorted(projects, key=lambda p: (not p.get("top"), p["id"])):
        L += [f'## {p["id"]}. {p["name"]}{"  ★ топ-10" if p.get("top") else ""}', "",
              f'**{p["material"]}** · сложность {p["difficulty"]}/5 · ~{p["hours"]} ч · **итого {p["total"]} ₽** (оснастка {p["tools_total"]} ₽)', "",
              p["summary"], "", f'Размер: {p.get("size","")}', "", "**Шаги:**"]
        L += [f"{i}. {s}" for i, s in enumerate(p["steps"], 1)]
        L += ["", "| Тип | Что купить | Кол-во | Цена | Сумма |", "|---|---|---|---|---|"]
        L += [f'| {l["kind"]} | [{l["title"]}]({l["url"]}) | {l["qty"]} {l["unit"]} | {l["price"]:.0f} | {l["sum"]:.0f} |' for l in p["bom"]]
        L += [f'| | **Итого** | | | **{p["total"]}** |', "", f'Источник идеи: {p["source"]}', ""]
    (ROOT / "out" / "PROJECTS.md").write_text("\n".join(L), encoding="utf-8")

if __name__ == "__main__":
    main()
